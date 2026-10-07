from fastapi import FastAPI
from fastapi.middleware.cors import CORSMiddleware
from pydantic import BaseModel
from openpyxl import Workbook, load_workbook
from datetime import datetime
from pathlib import Path
import threading

app = FastAPI()

# React frontend access
app.add_middleware(
    CORSMiddleware,
    allow_origins=["*"],
    allow_methods=["*"],
    allow_headers=["*"],
)

EXCEL_FILE = Path("PC_Specs.xlsx")

excel_lock = threading.Lock()


class PCSpecs(BaseModel):
    pc_name: str
    username: str
    windows: str
    cpu: str
    ram: str
    gpu: str
    disk: str
    manufacturer: str
    model: str
    ip_address: str


def create_excel():

    if not EXCEL_FILE.exists():

        wb = Workbook()
        ws = wb.active
        ws.title = "PC Specifications"

        headers = [
            "Date",
            "PC Name",
            "Username",
            "Windows",
            "CPU",
            "RAM",
            "GPU",
            "Disk",
            "Manufacturer",
            "Model",
            "IP Address"
        ]

        ws.append(headers)

        wb.save(EXCEL_FILE)


@app.get("/")
def home():

    return {
        "message": "PC Specification Server Running"
    }


@app.post("/save-pc")
def save_pc(data: PCSpecs):

    create_excel()

    with excel_lock:

        wb = load_workbook(EXCEL_FILE)
        ws = wb["PC Specifications"]

        ws.append([
            datetime.now().strftime("%Y-%m-%d %H:%M:%S"),
            data.pc_name,
            data.username,
            data.windows,
            data.cpu,
            data.ram,
            data.gpu,
            data.disk,
            data.manufacturer,
            data.model,
            data.ip_address
        ])

        wb.save(EXCEL_FILE)

    return {
        "success": True,
        "message": "PC information saved successfully"
    }